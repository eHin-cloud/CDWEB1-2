# -*- coding: utf-8 -*-
import sys
import docx
from lxml import etree

sys.stdout.reconfigure(encoding='utf-8')

def count_images_and_drawings(docx_path):
    print(f'=== File: {docx_path} ===')
    doc = docx.Document(docx_path)
    body = doc._body._element
    drawings = body.xpath('.//w:drawing | .//w:pict')
    print(f'Tong so drawings: {len(drawings)}')
    
    # Kiểm tra trong từng đoạn văn
    image_paras = []
    for i, p in enumerate(doc.paragraphs):
        p_xml = etree.tostring(p._p)
        if b'w:drawing' in p_xml or b'w:pict' in p_xml:
            next_txt = doc.paragraphs[i+1].text.strip() if i+1 < len(doc.paragraphs) else ''
            prev_txt = doc.paragraphs[i-1].text.strip() if i > 0 else ''
            image_paras.append((i, prev_txt, next_txt))
    print(f'So doan van chua hinh anh: {len(image_paras)}')
    for idx, prev_t, next_t in image_paras[:15]:
        print(f'  P[{idx}]: Prev="{prev_t[:45]}" | Next="{next_t[:45]}"')
    if len(image_paras) > 15:
        print(f'  ... va {len(image_paras) - 15} hinh nua')

count_images_and_drawings('BaoCao_GiaoDien_Message_30ChucNang (1).docx')
count_images_and_drawings('Báo cáo nhóm A.docx')
