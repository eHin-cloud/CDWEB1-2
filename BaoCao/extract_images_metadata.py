# -*- coding: utf-8 -*-
import sys
import docx
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path

sys.stdout.reconfigure(encoding='utf-8')

src_docx = 'BaoCao_GiaoDien_Message_30ChucNang (1).docx'
doc = docx.Document(src_docx)

# 1. Đọc relationships của word/document.xml
namespaces = {
    'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main',
    'r': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships',
    'a': 'http://schemas.openxmlformats.org/drawingml/2006/main',
    'wp': 'http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing',
    'pic': 'http://schemas.openxmlformats.org/drawingml/2006/picture'
}

with zipfile.ZipFile(src_docx, 'r') as z:
    rels_xml = z.read('word/_rels/document.xml.rels')
    rels_root = ET.fromstring(rels_xml)
    rId_to_target = {}
    for rel in rels_root:
        rId = rel.get('Id')
        target = rel.get('Target')
        if 'media/' in target:
            # target có dạng media/imageX.png
            rId_to_target[rId] = 'word/' + target if not target.startswith('word/') else target

print(f"Tổng số image relationships: {len(rId_to_target)}")

# 2. Duyệt các paragraph trong document
current_feat_num = 0
current_feat_name = ""
feat_images = {}

for p_idx, p in enumerate(doc.paragraphs):
    txt = p.text.strip()
    if txt.startswith('TÍNH NĂNG '):
        # Trích xuất số tính năng
        parts = txt.split(':')
        feat_head = parts[0].replace('TÍNH NĂNG', '').strip()
        try:
            current_feat_num = int(feat_head)
            current_feat_name = txt
            if current_feat_num not in feat_images:
                feat_images[current_feat_num] = []
        except:
            pass

    # Kiểm tra xem paragraph có chứa drawing không
    p_xml = p._p
    blips = p_xml.xpath('.//a:blip')
    if blips:
        for blip in blips:
            embed_id = blip.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed')
            media_path = rId_to_target.get(embed_id, 'UNKNOWN')
            
            # Tìm caption ở đoạn tiếp theo
            caption = ""
            desc = ""
            if p_idx + 1 < len(doc.paragraphs):
                p_next = doc.paragraphs[p_idx + 1].text.strip()
                if p_next.startswith('Hình '):
                    caption = p_next
                    if p_idx + 2 < len(doc.paragraphs):
                        p_desc = doc.paragraphs[p_idx + 2].text.strip()
                        if p_desc.startswith('Mô tả chức năng:'):
                            desc = p_desc
            
            feat_images[current_feat_num].append({
                'p_idx': p_idx,
                'embed_id': embed_id,
                'media_path': media_path,
                'caption': caption,
                'desc': desc
            })

print("\n=== THỐNG KÊ HÌNH ẢNH THEO TỪNG TÍNH NĂNG (1-30) ===")
total_imgs = 0
for f_num in range(1, 31):
    imgs = feat_images.get(f_num, [])
    total_imgs += len(imgs)
    print(f"Tính năng {f_num:2d}: có {len(imgs)} hình ảnh")
    for img in imgs:
        print(f"    - Media: {img['media_path']} | Caption: {img['caption'][:50]}")

print(f"\nTổng số hình ảnh tìm thấy trên 30 tính năng: {total_imgs}")
