import os
import sys
import re
from copy import deepcopy
import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn
from lxml import etree

sys.stdout.reconfigure(encoding='utf-8')

def set_cell_shading(cell, color_hex="F1F5F9"):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    tcPr.append(shd)

def set_table_borders(table, color="CBD5E1", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'  <w:top w:w="{top}" w:type="dxa"/>'
        f'  <w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'  <w:left w:w="{left}" w:type="dxa"/>'
        f'  <w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def format_run(run, font_name="Times New Roman", size_pt=13, bold=False, italic=False, color=None):
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = color
    # Ensure XML has w:rFonts
    rPr = run._r.get_or_add_rPr()
    rFonts = parse_xml(f'<w:rFonts {nsdecls("w")} w:ascii="{font_name}" w:hAnsi="{font_name}" w:cs="{font_name}"/>')
    rPr.append(rFonts)

def map_caption_to_feat(caption):
    cap = caption.lower()
    if 'input lỗi' in cap: return 'FEAT-HIEN-03'
    if 'webauthn' in cap or 'passkey' in cap: return 'FEAT-HIEN-04'
    if 'ký số' in cap or 'pdf hợp đồng' in cap: return 'FEAT-HIEN-06'
    if 'kiểm duyệt hồ sơ' in cap or 'bảng điều khiển kiểm duyệt' in cap: return 'FEAT-HIEN-07'
    if 'nhật ký kiểm toán' in cap: return 'FEAT-HIEN-07'
    if 'onboarding' in cap or 'thiết lập cơ sở lưu trú ban đầu' in cap: return 'FEAT-HIEN-09'
    if 'quản lý và thiết lập cơ sở' in cap or 'quản lý cơ sở' in cap: return 'FEAT-QUY-01'
    if 'thêm mới và cập nhật thông tin phòng' in cap or 'thêm / cập nhật thông tin phòng' in cap: return 'FEAT-QUY-02'
    if 'sơ đồ ma trận' in cap or 'ma trận phòng' in cap: return 'FEAT-QUY-03'
    if 'chốt số điện - nước' in cap or 'chốt số điện nước' in cap: return 'FEAT-QUY-04'
    if 'bảng kê thanh toán' in cap or 'quản lý thanh toán' in cap or 'vietqr' in cap: return 'FEAT-QUY-07'
    if 'quản lý trang thiết bị' in cap or 'quản lý tài sản' in cap: return 'FEAT-QUY-08'
    if 'sổ quỹ' in cap: return 'FEAT-QUY-09'
    if 'cổng tìm kiếm phòng' in cap or 'cổng renty' in cap or 'review lưu trú' in cap: return 'FEAT-VINHEM-01'
    if 'trợ lý ảo ai' in cap or 'chatbot tư vấn' in cap: return 'FEAT-VINHEM-04'
    if 'cổng thông tin cư dân' in cap or 'cổng dịch vụ cư dân' in cap or 'guest portal' in cap: return 'FEAT-VINHEM-05'
    if 'tiếp nhận sự cố' in cap or 'smart ticket' in cap: return 'FEAT-VINHEM-06'
    if 'tờ khai thay đổi thông tin cư trú' in cap or 'mẫu ct01' in cap or 'tờ khai ct01' in cap: return 'FEAT-VINHEM-08'
    return 'UNMAPPED'

def main():
    print("Bắt đầu xử lý đồng bộ đặc tả kỹ thuật Enterprise vào Word Document...")
    
    # 1. Load backup document
    backup_path = 'BaoCao/BaoCao_NhomA_backup.docx'
    enterprise_output_path = 'BaoCao/BaoCao_NhomA_Enterprise.docx'
    official_output_path = 'BaoCao/BaoCao_NhomA.docx'
    
    doc = docx.Document(backup_path)
    body = doc._body._element
    children = list(body)
    
    # Locate Section IV and References in backup body (ignore TOC sdt at index 30)
    sec4_start_idx = -1
    sec4_end_idx = -1
    for idx, c in enumerate(children):
        if idx > 50 and c.tag.endswith('p'):
            txt = ''.join(c.itertext()).strip()
            if ('IV. THIẾT KẾ' in txt or 'IV. ĐẶC TẢ' in txt) and sec4_start_idx == -1:
                sec4_start_idx = idx
            if 'TÀI LIỆU THAM KHẢO' in txt and sec4_start_idx != -1 and sec4_end_idx == -1:
                sec4_end_idx = idx
                break
            
    print(f"Section IV vị trí trong backup docx: start={sec4_start_idx}, end={sec4_end_idx}")
    
    # 2. Extract and preserve all 32 drawings + captions
    images_by_feat = {}
    for idx in range(sec4_start_idx, sec4_end_idx):
        c = children[idx]
        xml_bytes = etree.tostring(c)
        if b'w:drawing' in xml_bytes or b'w:pict' in xml_bytes:
            next_c = children[idx+1] if idx+1 < len(children) else None
            caption = ''.join(next_c.itertext()).strip() if next_c is not None else ''
            feat = map_caption_to_feat(caption)
            if feat not in images_by_feat:
                images_by_feat[feat] = []
            images_by_feat[feat].append((deepcopy(c), deepcopy(next_c), caption))
            
    total_images_saved = sum(len(v) for v in images_by_feat.values())
    print(f"Đã trích xuất và bảo toàn toàn bộ {total_images_saved} hình ảnh/wireframe từ Section IV.")
    for feat, imgs in sorted(images_by_feat.items()):
        print(f"  - {feat}: {len(imgs)} ảnh")

    # 3. Reference target element where new Section IV will be inserted before
    target_ref_element = children[sec4_end_idx]
    
    # Remove old Section IV elements from body
    for idx in range(sec4_end_idx - 1, sec4_start_idx - 1, -1):
        body.remove(children[idx])
        
    print(f"Đã dọn dẹp {sec4_end_idx - sec4_start_idx} phần tử cũ của Section IV.")

    # 4. Read BaoCao_NhomA_Moi.md
    with open('BaoCao/BaoCao_NhomA_Moi.md', 'r', encoding='utf-8') as f:
        md_text = f.read()
        
    sec4_md_start = md_text.find('# IV. ĐẶC TẢ KỸ THUẬT VÀ KỊCH BẢN XỬ LÝ LỖI')
    sec4_md_end = md_text.find('# TÀI LIỆU THAM KHẢO')
    sec4_md = md_text[sec4_md_start:sec4_md_end].strip()
    
    lines = sec4_md.split('\n')
    print(f"Đọc {len(lines)} dòng markdown từ Section IV.")
    
    # Helper doc for creating elements
    helper_doc = docx.Document()
    
    current_feat_code = None
    images_inserted_for_feat = set()
    
    def insert_element(elem):
        target_ref_element.addprevious(deepcopy(elem))
        
    def add_p(text, size_pt=13, bold=False, italic=False, align=WD_ALIGN_PARAGRAPH.LEFT, space_before=2, space_after=4):
        p = helper_doc.add_paragraph()
        p.alignment = align
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = 1.15
        run = p.add_run(text)
        format_run(run, size_pt=size_pt, bold=bold, italic=italic)
        insert_element(p._p)
        return p

    def add_code_block(code_text):
        p = helper_doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.2)
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(code_text)
        format_run(run, font_name="Consolas", size_pt=10, color=RGBColor(30, 41, 59))
        pPr = p._p.get_or_add_pPr()
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F8FAFC"/>')
        pPr.append(shd)
        insert_element(p._p)

    def set_cell_width(cell, width_dxa):
        tcPr = cell._tc.get_or_add_tcPr()
        tcW = parse_xml(f'<w:tcW {nsdecls("w")} w:w="{width_dxa}" w:type="dxa"/>')
        tcPr.append(tcW)

    def add_table_from_rows(table_rows):
        if not table_rows or len(table_rows) < 2:
            return
        headers = table_rows[0]
        data_rows = table_rows[1:]
        
        num_cols = len(headers)
        tbl = helper_doc.add_table(rows=len(data_rows) + 1, cols=num_cols)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders(tbl, color="CBD5E1", sz="4")

        # Determine column widths (Total ~9360 dxa for 6.5in printable width)
        total_dxa = 9360
        hdr_0 = headers[0].lower()
        if num_cols == 2:
            if 'nguyên nhân' in hdr_0 or 'thao tác' in hdr_0:
                col_widths = [3300, 6060]  # 35% vs 65% for Edge Cases table
            else:
                col_widths = [2800, 6560]  # 30% vs 70% for Metadata
        elif num_cols == 3:
            col_widths = [2800, 3280, 3280]
        elif num_cols == 4:
            if 'method' in hdr_0:
                col_widths = [1200, 2760, 2200, 3200]
            elif 'tình huống' in hdr_0:
                col_widths = [2800, 1160, 2400, 3000]
            else:
                col_widths = [total_dxa // num_cols] * num_cols
        elif num_cols == 5:
            col_widths = [1700, 1400, 1160, 2550, 2550]
        else:
            col_widths = [total_dxa // num_cols] * num_cols
        
        # Header row
        hdr_row = tbl.rows[0]
        trPr = hdr_row._tr.get_or_add_trPr()
        trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
        
        for c_idx, h_text in enumerate(headers):
            cell = hdr_row.cells[c_idx]
            cell.text = h_text.strip()
            set_cell_width(cell, col_widths[c_idx] if c_idx < len(col_widths) else 1500)
            set_cell_shading(cell, "E2E8F0")
            set_cell_margins(cell, top=120, bottom=120, left=150, right=150)
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)
                for r in p.runs:
                    format_run(r, size_pt=11.5, bold=True, color=RGBColor(15, 23, 42))
                    
        # Data rows
        for r_idx, d_row in enumerate(data_rows):
            row = tbl.rows[r_idx + 1]
            bg_color = "FFFFFF" if r_idx % 2 == 0 else "F8FAFC"
            for c_idx in range(num_cols):
                cell = row.cells[c_idx]
                val = d_row[c_idx].strip() if c_idx < len(d_row) else ""
                val_clean = val.replace('<br>', '\n').replace('\\|', '|').replace('`', '')
                cell.text = val_clean
                set_cell_width(cell, col_widths[c_idx] if c_idx < len(col_widths) else 1500)
                set_cell_shading(cell, bg_color)
                set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
                for p in cell.paragraphs:
                    p.paragraph_format.space_before = Pt(2)
                    p.paragraph_format.space_after = Pt(2)
                    for r in p.runs:
                        format_run(r, size_pt=11)
                        
        insert_element(tbl._tbl)

    def check_and_insert_feature_images(feat_code):
        if feat_code in images_by_feat and feat_code not in images_inserted_for_feat:
            for img_el, cap_el, cap_text in images_by_feat[feat_code]:
                insert_element(img_el)
                insert_element(cap_el)
            images_inserted_for_feat.add(feat_code)

    # 5. Parse and Render Markdown
    i = 0
    in_code_block = False
    code_buffer = []
    
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()
        
        # Check code block
        if stripped.startswith('```'):
            if in_code_block:
                in_code_block = False
                add_code_block('\n'.join(code_buffer))
                code_buffer = []
            else:
                in_code_block = True
                code_buffer = []
            i += 1
            continue
            
        if in_code_block:
            code_buffer.append(line)
            i += 1
            continue
            
        # Check table
        if stripped.startswith('|') and stripped.endswith('|'):
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith('|') and lines[i].strip().endswith('|'):
                table_lines.append(lines[i].strip())
                i += 1
            parsed_rows = []
            for tl in table_lines:
                # ignore separator row like | :--- | :--- |
                cols = [c.strip() for c in tl.split('|')[1:-1]]
                if all(re.match(r'^:?-+:?$', c) for c in cols if c):
                    continue
                parsed_rows.append(cols)
            if parsed_rows:
                add_table_from_rows(parsed_rows)
            continue
            
        # Check Headings
        if stripped.startswith('# '):
            add_p(stripped[2:].strip(), size_pt=16, bold=True, space_before=14, space_after=6)
        elif stripped.startswith('## '):
            add_p(stripped[3:].strip(), size_pt=14.5, bold=True, space_before=12, space_after=6)
        elif stripped.startswith('### '):
            heading_txt = stripped[4:].strip()
            add_p(heading_txt, size_pt=13.5, bold=True, space_before=10, space_after=4)
        elif stripped.startswith('#### '):
            heading_txt = stripped[5:].strip()
            add_p(heading_txt, size_pt=13, bold=True, space_before=8, space_after=3)
            # Check if this is section 3 (Hợp đồng dữ liệu), insert images if under section 2
            if '3. Hợp đồng Dữ liệu' in heading_txt and current_feat_code:
                check_and_insert_feature_images(current_feat_code)
        elif stripped.startswith('##### '):
            add_p(stripped[6:].strip(), size_pt=13, bold=True, italic=True, space_before=6, space_after=2)
        elif stripped.startswith('- [ ]') or stripped.startswith('- [x]'):
            p = helper_doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.25)
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            run_box = p.add_run('☑ ' if '[x]' in stripped else '☐ ')
            format_run(run_box, size_pt=13, bold=True, color=RGBColor(37, 99, 235))
            run_txt = p.add_run(stripped[5:].strip())
            format_run(run_txt, size_pt=13)
            insert_element(p._p)
        elif stripped.startswith('- ') or stripped.startswith('* '):
            p = helper_doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.25)
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            run_bullet = p.add_run('• ')
            format_run(run_bullet, size_pt=13, bold=True)
            run_txt = p.add_run(stripped[2:].strip())
            format_run(run_txt, size_pt=13)
            insert_element(p._p)
        elif stripped == '---':
            # End of feature, ensure images are inserted
            if current_feat_code:
                check_and_insert_feature_images(current_feat_code)
                current_feat_code = None
        elif stripped:
            # Check if line mentions feature code
            feat_match = re.search(r'FEAT-[A-Z]+-\d+', stripped)
            if feat_match:
                current_feat_code = feat_match.group(0)
            add_p(stripped, size_pt=13)
            
        i += 1

    # Insert any remaining feature images
    for feat_code in images_by_feat:
        check_and_insert_feature_images(feat_code)

    # 6. Global Typography Formatting Pass (Times New Roman 100%, 13pt body)
    print("Thực hiện định dạng chuẩn hóa phông chữ Times New Roman và cỡ chữ toàn văn bản...")
    for p in doc.paragraphs:
        txt = p.text.strip()
        is_heading = any(p.style.name.startswith(h) for h in ['Heading', 'Title']) or \
                     txt.startswith(('I.', 'II.', 'III.', 'IV.', 'A.', 'B.', 'C.', 'CHƯƠNG', 'ĐẶC TẢ', 'TRƯỜNG', 'BÁO CÁO'))
        for r in p.runs:
            format_run(r, font_name="Times New Roman", size_pt=14 if is_heading else 13, bold=r.bold, italic=r.italic)

    for tbl in doc.tables:
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        for r_idx, row in enumerate(tbl.rows):
            for cell in row.cells:
                for p in cell.paragraphs:
                    for r in p.runs:
                        format_run(r, font_name="Times New Roman", size_pt=11.5 if r_idx == 0 else 11, bold=r.bold, italic=r.italic)

    # 7. Save to Enterprise docx
    doc.save(enterprise_output_path)
    print(f"Đã lưu thành công văn bản hoàn chỉnh vào: {enterprise_output_path}")

    # 8. Save/copy to official BaoCao_NhomA.docx
    try:
        doc.save(official_output_path)
        print(f"Đã lưu thành công văn bản chính thức vào: {official_output_path}")
    except Exception as e:
        print(f"Cảnh báo: Không thể ghi đè trực tiếp {official_output_path} ({e}). File đã sẵn sàng tại {enterprise_output_path}")

if __name__ == '__main__':
    main()
